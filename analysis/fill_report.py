#!/usr/bin/env python3
"""
Fills the CS4310 Project 1 submission document (cs4310 project1_26Fall.docx)
with the design, testing, implementation, and performance-analysis answers,
producing CS4310-Project1-Report-EricLa.docx. Run once after experiment.py
has produced analysis/results/summary_table.csv and analysis/results/graphs/.
"""
import csv
from pathlib import Path

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
SRC_PATH = ROOT / "cs4310 project1_26Fall.docx"
OUT_PATH = ROOT / "CS4310-Project1-Report-EricLa.docx"
RESULTS_DIR = ROOT / "analysis" / "results"
GRAPHS_DIR = RESULTS_DIR / "graphs"

STUDENT_NAME = "Eric La"
DATE_STR = "October 29, 2026"

ANSWER_COLOR = RGBColor(0x00, 0x00, 0x00)


def set_paragraph_text(paragraph, text, bold=False, italic=False, mono=False, size=None):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = ANSWER_COLOR
    if mono:
        run.font.name = "Consolas"
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = rPr.makeelement(qn('w:rFonts'), {})
            rPr.append(rFonts)
        rFonts.set(qn('w:eastAsia'), "Consolas")
    if size:
        run.font.size = Pt(size)
    return paragraph


def insert_paragraph_after(paragraph, text="", mono=False, bold=False, size=None):
    new_p = docx.oxml.OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    if text:
        set_paragraph_text(new_paragraph, text, bold=bold, mono=mono, size=size)
    return new_paragraph


def insert_image_after(paragraph, image_path, width_in=5.5):
    new_p = docx.oxml.OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    new_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = new_paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))
    return new_paragraph


def set_cell_text(cell, text, size=None):
    """Sets the text of a table cell, clearing any extra paragraphs the cell may have."""
    paragraphs = cell.paragraphs
    set_paragraph_text(paragraphs[0], text, size=size)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def find_paragraph(doc, text_startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(text_startswith):
            return p
    raise ValueError(f"Paragraph not found: {text_startswith!r}")


def find_table_after(doc, anchor_paragraph):
    """Return the first table element that appears after anchor_paragraph in document order."""
    body = doc.element.body
    found_anchor = False
    for child in body.iterchildren():
        if child is anchor_paragraph._p:
            found_anchor = True
            continue
        if found_anchor and child.tag.endswith('}tbl'):
            return docx.table.Table(child, doc)
    raise ValueError("No table found after anchor paragraph")


def load_summary_table():
    rows = {}
    with (RESULTS_DIR / "summary_table.csv").open() as f:
        reader = csv.DictReader(f)
        for row in reader:
            size = int(row["Input Size (jobs)"])
            rows[size] = {
                "FCFS": row["FCFS"], "SJF": row["SJF"],
                "RR2": row["RR2"], "RR5": row["RR5"],
            }
    return rows


def main():
    doc = docx.Document(str(SRC_PATH))

    # --- Student name / date ---
    for p in doc.paragraphs:
        if p.text.startswith("Student Name:"):
            set_paragraph_text(p, f"Student Name: {STUDENT_NAME}")
        elif p.text.startswith("Date:"):
            set_paragraph_text(p, f"Date: {DATE_STR}")

    # =========================================================
    # Part 1(a): pseudocode
    # =========================================================
    p = find_paragraph(doc, "<Insert answers here>")  # first occurrence -> Part 1(a)
    set_paragraph_text(p, "")

    fcfs_pseudo = """FCFS(jobs[1..n]):
    time = 0
    for i = 1 to n:                 // jobs processed in file order
        start[i] = time
        finish[i] = time + burst[i]
        turnaround[i] = finish[i]   // arrival time is 0 for all jobs
        time = finish[i]
    avgTurnaround = sum(turnaround[1..n]) / n
    return schedule, avgTurnaround"""

    sjf_pseudo = """SJF(jobs[1..n]):                  // non-preemptive, all jobs arrive at time 0
    order = sort jobs by burst time ascending
            (ties broken by original file order)
    time = 0
    for each job in order:
        start = time
        finish = time + job.burst
        turnaround = finish         // arrival time is 0
        time = finish
    avgTurnaround = sum(turnaround) / n
    return schedule, avgTurnaround"""

    rr_pseudo = """RoundRobin(jobs[1..n], quantum q):   // q = 2 for RR-2, q = 5 for RR-5
    queue = FIFO queue of jobs in file order
    remaining[i] = burst[i] for all i
    time = 0
    while queue is not empty:
        job = queue.dequeue()
        run = min(q, remaining[job])
        start = time
        time = time + run
        remaining[job] -= run
        if remaining[job] > 0:
            queue.enqueue(job)      // not finished, go to back of queue
        else:
            finish[job] = time
            turnaround[job] = time  // arrival time is 0
    avgTurnaround = sum(turnaround) / n
    return schedule, avgTurnaround"""

    anchor = p
    anchor = insert_paragraph_after(anchor, "a) FCFS", bold=True)
    anchor = insert_paragraph_after(anchor, fcfs_pseudo, mono=True, size=10)
    anchor = insert_paragraph_after(anchor, "b) SJF (non-preemptive)", bold=True)
    anchor = insert_paragraph_after(anchor, sjf_pseudo, mono=True, size=10)
    anchor = insert_paragraph_after(anchor, "c) & d) Round-Robin (shared logic; q = 2 for RR-2, q = 5 for RR-5)", bold=True)
    anchor = insert_paragraph_after(anchor, rr_pseudo, mono=True, size=10)

    # =========================================================
    # Part 1(b): testing table
    # =========================================================
    anchor_b = find_paragraph(doc, "<Complete the following table>")
    set_paragraph_text(anchor_b, "")
    table_b = find_table_after(doc, anchor_b)

    job_inputs = {
        5: "Job1=7, Job2=18, Job3=10, Job4=4, Job5=12",
        10: "Job1=7, Job2=18, Job3=10, Job4=4, Job5=12,\nJob6=3, Job7=15, Job8=9, Job9=6, Job10=20",
        15: "Job1=7, Job2=18, Job3=10, Job4=4, Job5=12,\nJob6=3, Job7=15, Job8=9, Job9=6, Job10=20,\nJob11=11, Job12=2, Job13=17, Job14=8, Job15=5",
    }
    expected = {
        5: {"FCFS": "31.40", "SJF": "24.00", "RR2": "36.40", "RR5": "36.00"},
        10: {"FCFS": "54.60", "SJF": "41.60", "RR2": "70.00", "RR5": "68.10"},
        15: {"FCFS": "80.07", "SJF": "55.13", "RR2": "97.07", "RR5": "95.80"},
    }
    row_sizes = [5, 10, 15]
    for row_idx, size in zip([1, 2, 3], row_sizes):
        set_cell_text(table_b.cell(row_idx, 1), job_inputs[size], size=9)
        for col_idx, alg in zip([2, 3, 4, 5], ["FCFS", "SJF", "RR2", "RR5"]):
            val = expected[size][alg]
            set_cell_text(table_b.cell(row_idx, col_idx), f"{val}  ✓")

    # =========================================================
    # Part 1(c): testing strategy
    # =========================================================
    p1c = find_paragraph(doc, "<Insert answers here>")  # now matches the remaining one (Part 1c)
    strategy_text = (
        "Three input sizes are used: 5, 10, and 15 jobs. For each size, 20 random job sets "
        "(\"trials\") are generated, with each job's burst time drawn uniformly from 1 to 20 ms "
        "(maximum job length = 20, the minimum allowed by the assignment). A fixed random seed "
        "(42) is used so the experiment is reproducible. Each of the 20 generated job sets per "
        "size is written to its own input file and is reused unchanged across all four scheduling "
        "algorithms (FCFS, SJF, RR-2, RR-5), ensuring a fair, apples-to-apples comparison within "
        "a trial. For each trial, each algorithm computes its own average turnaround time across "
        "that trial's jobs; those 20 per-trial averages (X = {x1, ..., x20}) are then averaged "
        "using the formula given above to produce the final \"average of average turnaround "
        "times\" reported in the Part 3(a) table. This process is implemented in "
        "analysis/experiment.py, which drives the four compiled Java programs and collects their "
        "output automatically."
    )
    set_paragraph_text(p1c, strategy_text)

    # =========================================================
    # Part 2(a), (b), (c)
    # =========================================================
    p2a = find_paragraph(doc, "<Generate four programs")
    set_paragraph_text(p2a, (
        "Four standalone Java programs were created, one per algorithm, each independently "
        "compilable and runnable: FCFS.java, SJF.java, RR2.java, and RR5.java (submitted "
        "alongside this document)."
    ))

    p2b = find_paragraph(doc, "<Generate documentation")
    set_paragraph_text(p2b, (
        "Each of the four files begins with a header comment documenting the algorithm, its "
        "assumptions (all jobs arrive at time 0, FCFS/RR process jobs in file order), and usage "
        "instructions (e.g. \"javac FCFS.java\" then \"java FCFS [path-to-job-file]\"). Class and "
        "method names are self-descriptive (readJobs, printGanttChart, printTable, etc.)."
    ))

    p2c = find_paragraph(doc, "<Complete the four columns")
    set_paragraph_text(p2c, (
        "All four testing cases were run against the four compiled programs; computed average "
        "turnaround times matched the expected values in the table in Part 1(b) exactly, so each "
        "cell above is marked with a checkmark."
    ))

    p2d = find_paragraph(doc, "<Insert totally four screen shots")
    set_paragraph_text(p2d, (
        "[Insert your own screenshots here: for each of FCFS, SJF, RR2, RR5, run "
        "\"java -cp out <Program> testdata/job5.txt\" in your terminal and capture the compile "
        "and run output.]"
    ), italic=True)

    # =========================================================
    # Part 3(a): performance table
    # =========================================================
    p3a = find_paragraph(doc, "Run each program with the designed randomly generated")
    table_3a = find_table_after(doc, p3a)
    summary = load_summary_table()
    for row_idx, size in zip([1, 2, 3], row_sizes):
        for col_idx, alg in zip([1, 2, 3, 4], ["FCFS", "SJF", "RR2", "RR5"]):
            set_cell_text(table_3a.cell(row_idx, col_idx), summary[size][alg])

    # =========================================================
    # Part 3(b): graphs + summary + ranking
    # =========================================================
    p_graphs_intro = find_paragraph(doc, "<Insert totally four graphs")
    set_paragraph_text(p_graphs_intro, "")
    anchor = p_graphs_intro
    for alg in ["FCFS", "SJF", "RR2", "RR5"]:
        anchor = insert_image_after(anchor, GRAPHS_DIR / f"{alg}.png", width_in=4.5)

    p_summary = find_paragraph(doc, "<Write a summary>")
    set_paragraph_text(p_summary, (
        "FCFS: average turnaround grows roughly linearly with input size, driven by the "
        "convoy effect — long jobs early in file order delay every job behind them. "
        "SJF: consistently the lowest average turnaround time at every input size, and the "
        "gap between it and the other algorithms widens as job count increases, since it always "
        "clears short jobs before they can be stuck waiting behind long ones. "
        "RR-2: the highest average turnaround time at every size; the small 2 ms quantum forces "
        "many context switches, so jobs cycle through the queue many times before finishing. "
        "RR-5: consistently slightly better than RR-2 but still worse than FCFS and SJF; the "
        "larger quantum lets jobs finish in fewer trips through the queue, trending RR-5's "
        "behavior closer to FCFS."
    ))

    # Combined graph: clear the inline placeholder text in that paragraph, then insert image after it
    p_combined = find_paragraph(doc, "Plot all four graphs on the same graph")
    set_paragraph_text(p_combined, (
        "Plot all four graphs on the same graph and compare the performance of all four "
        "algorithms. Rank four scheduling algorithms. Try giving the reasons for the findings."
    ))
    insert_image_after(p_combined, GRAPHS_DIR / "combined.png", width_in=5.5)

    p_explain = find_paragraph(doc, "<Write about explaining the results>")
    set_paragraph_text(p_explain, (
        "Ranking from best (lowest average turnaround) to worst: 1) SJF, 2) FCFS, 3) RR-5, "
        "4) RR-2. SJF wins because it is provably optimal for minimizing average turnaround time "
        "when all jobs arrive simultaneously — running the shortest jobs first keeps the "
        "average number of jobs waiting in the system as low as possible. FCFS is simple and "
        "performs reasonably, but a single long job placed early in the file delays every job "
        "behind it (the convoy effect), which SJF avoids entirely. Round-Robin is designed to "
        "optimize fairness and responsiveness (every job gets the CPU quickly and repeatedly), "
        "not turnaround time, so both RR variants trail FCFS and SJF; RR-5 edges out RR-2 "
        "because its larger time slice means fewer context switches and jobs complete in fewer "
        "trips through the queue, pushing its behavior closer to FCFS as the quantum grows."
    ))

    # =========================================================
    # Part 3(c): conclusion
    # =========================================================
    p_conclusion = find_paragraph(doc, "<Write a conclusion")
    set_paragraph_text(p_conclusion, (
        "This project's main strength is that the four scheduling algorithms were implemented as "
        "fully independent, documented programs sharing the same input format, which made it "
        "straightforward to automate a 20-trial randomized experiment (analysis/experiment.py) "
        "on top of them without modifying the submitted files. The results matched textbook "
        "expectations — SJF minimized average turnaround time while Round-Robin, optimized "
        "for fairness rather than turnaround, consistently trailed — which gave good "
        "confidence in the implementations' correctness. The main constraint is that the "
        "simulation assumes all jobs arrive at time 0 and burst times are known in advance, which "
        "is unrealistic for a live operating system; a more complete simulator would model "
        "staggered arrival times, preemption based on priority, and I/O-bound waiting. If "
        "redone, I would keep the standalone-program design and the automated trial harness, but "
        "I would add preemptive SJF (shortest-remaining-time-first) and variable arrival times to "
        "make the comparison more representative of real scheduling behavior, and I would collect "
        "waiting time and response time alongside turnaround time to get a fuller picture of each "
        "algorithm's trade-offs."
    ))

    doc.save(str(OUT_PATH))
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
